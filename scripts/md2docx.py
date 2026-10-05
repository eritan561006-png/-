"""将实施方案 Markdown 按公文排版规范转换为 Word 文档。

用法：python3 scripts/md2docx.py <输入.md> <输出.docx>
依赖：pip install python-docx
"""
import re
import sys

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

FONT_TITLE = "方正小标宋简体"
FONT_HEI = "黑体"
FONT_KAI = "楷体"
FONT_FANG = "仿宋"
SIZE_BODY = Pt(16)   # 三号
SIZE_TABLE = Pt(10.5)  # 五号


def set_font(run, name, size=None, bold=None, color=None):
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rfonts.set(qn(attr), name)
    if size is not None:
        run.font.size = size
    if bold is not None:
        run.font.bold = bold
    if color is not None:
        run.font.color.rgb = color


def para_format(p, first_indent=True, line=Pt(28), before=0, after=0, align=None):
    pf = p.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    pf.line_spacing = line
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    if first_indent:
        pf.first_line_indent = Pt(32)
    if align is not None:
        p.alignment = align


def add_inline(p, text, font, size, bold=False):
    """支持 **加粗** 的行内文本。"""
    parts = re.split(r"(\*\*[^*]+\*\*)", text)
    for part in parts:
        if not part:
            continue
        is_bold = part.startswith("**") and part.endswith("**")
        content = part[2:-2] if is_bold else part
        content = content.replace("`", "")
        run = p.add_run(content)
        set_font(run, FONT_HEI if (is_bold and font == FONT_FANG) else font, size, bold or is_bold)


def shade(cell, fill):
    tcpr = cell._element.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tcpr.append(shd)


def add_field(paragraph, instr):
    run = paragraph.add_run()
    for tag, extra in (("begin", None), ("instr", instr), ("separate", None), ("end", None)):
        if tag == "instr":
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = extra
        else:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), tag)
        run._element.append(el)
    return run


def repeat_header(row):
    trpr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    trpr.append(el)


def add_table(doc, rows):
    header, body = rows[0], rows[1:]
    ncols = len(header)
    table = doc.add_table(rows=1 + len(body), cols=ncols)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for r_idx, row in enumerate([header] + body):
        cells = table.rows[r_idx].cells
        for c_idx in range(ncols):
            text = row[c_idx] if c_idx < len(row) else ""
            cell = cells[c_idx]
            cell.text = ""
            p = cell.paragraphs[0]
            para_format(p, first_indent=False, line=Pt(18))
            if r_idx == 0:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                add_inline(p, text, FONT_HEI, SIZE_TABLE, bold=True)
                shade(cell, "D9E2F3")
            else:
                add_inline(p, text, FONT_FANG, SIZE_TABLE)
    repeat_header(table.rows[0])
    # 首列为空时与上一行纵向合并（用于分组表格）
    for r_idx in range(len(body), 0, -1):
        if not body[r_idx - 1][0].strip() and r_idx > 1:
            top = r_idx - 1
            while top > 1 and not body[top - 1][0].strip():
                top -= 1
            a = table.cell(top, 0)
            b = table.cell(r_idx, 0)
            if a._tc is not b._tc:
                a.merge(b)
    doc.add_paragraph()


def add_code(doc, lines):
    table = doc.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    cell = table.cell(0, 0)
    shade(cell, "F2F2F2")
    cell.text = ""
    for i, line in enumerate(lines):
        p = cell.paragraphs[0] if i == 0 else cell.add_paragraph()
        para_format(p, first_indent=False, line=Pt(20))
        run = p.add_run(line)
        set_font(run, FONT_KAI, Pt(12))
    doc.add_paragraph()


def build_cover(doc, title, meta):
    for _ in range(6):
        doc.add_paragraph()
    p = doc.add_paragraph()
    para_format(p, first_indent=False, line=Pt(44), align=WD_ALIGN_PARAGRAPH.CENTER)
    head, _, tail = title.partition("“")
    if tail:
        set_font(p.add_run(head), FONT_TITLE, Pt(26))
        p.add_run().add_break()
        set_font(p.add_run("“" + tail), FONT_TITLE, Pt(26))
    else:
        set_font(p.add_run(title), FONT_TITLE, Pt(26))
    for _ in range(8):
        doc.add_paragraph()
    for k, v in meta:
        if k in ("编制单位", "文档版本", "编制日期"):
            p = doc.add_paragraph()
            para_format(p, first_indent=False, line=Pt(32), align=WD_ALIGN_PARAGRAPH.CENTER)
            set_font(p.add_run(f"{k}：{v}"), FONT_KAI, SIZE_BODY)
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    # 目录页
    p = doc.add_paragraph()
    para_format(p, first_indent=False, line=Pt(40), align=WD_ALIGN_PARAGRAPH.CENTER)
    set_font(p.add_run("目　录"), FONT_HEI, Pt(22))
    p = doc.add_paragraph()
    add_field(p, 'TOC \\o "1-2" \\h \\z \\u')
    note = doc.add_paragraph()
    set_font(note.add_run("（打开文档后右键此处选择“更新域”即可生成目录）"), FONT_KAI, Pt(12),
             color=RGBColor(0x80, 0x80, 0x80))


def setup_document():
    doc = Document()
    sec = doc.sections[0]
    sec.page_height, sec.page_width = Cm(29.7), Cm(21.0)
    sec.top_margin, sec.bottom_margin = Cm(3.7), Cm(3.5)
    sec.left_margin, sec.right_margin = Cm(2.8), Cm(2.6)
    normal = doc.styles["Normal"]
    normal.font.name = FONT_FANG
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), FONT_FANG)
    normal.font.size = SIZE_BODY
    for name, font, size in (("Heading 1", FONT_HEI, SIZE_BODY), ("Heading 2", FONT_KAI, SIZE_BODY),
                             ("Heading 3", FONT_FANG, SIZE_BODY), ("Heading 4", FONT_FANG, SIZE_BODY)):
        st = doc.styles[name]
        st.font.name = font
        st.font.size = size
        st.font.color.rgb = RGBColor(0, 0, 0)
        st.font.bold = name in ("Heading 1", "Heading 3")
        st.element.rPr.rFonts.set(qn("w:eastAsia"), font)
    return doc


def add_footer_page_number(section):
    footer = section.footer
    footer.is_linked_to_previous = False
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(p.add_run("— "), FONT_FANG, Pt(14))
    set_font(add_field(p, "PAGE"), FONT_FANG, Pt(14))
    set_font(p.add_run(" —"), FONT_FANG, Pt(14))


def convert(src, dst):
    lines = open(src, encoding="utf-8").read().splitlines()
    doc = setup_document()

    title = lines[0].lstrip("# ").strip()
    meta, i = [], 1
    while i < len(lines) and (lines[i].startswith(">") or not lines[i].strip()):
        if lines[i].startswith(">"):
            k, _, v = lines[i].lstrip("> ").partition("：")
            meta.append((k, v))
        i += 1
    build_cover(doc, title, meta)

    body = doc.add_section(WD_SECTION.NEW_PAGE)
    add_footer_page_number(body)
    p = doc.add_paragraph()
    para_format(p, first_indent=False, line=Pt(36), after=12, align=WD_ALIGN_PARAGRAPH.CENTER)
    set_font(p.add_run(title), FONT_TITLE, Pt(22))
    for k, v in meta:
        if k in ("适用范围", "说明"):
            p = doc.add_paragraph()
            para_format(p)
            add_inline(p, f"{k}：{v}", FONT_KAI, SIZE_BODY)

    while i < len(lines):
        line = lines[i]
        s = line.strip()
        if not s:
            i += 1
            continue
        if s.startswith("```"):
            block, i = [], i + 1
            while i < len(lines) and not lines[i].strip().startswith("```"):
                block.append(lines[i])
                i += 1
            add_code(doc, block)
            i += 1
            continue
        if s.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-{3,}:?", c) for c in cells):
                    rows.append(cells)
                i += 1
            add_table(doc, rows)
            continue
        m = re.match(r"^(#{2,5})\s+(.*)$", s)
        if m:
            level = len(m.group(1)) - 1
            text = m.group(2)
            if level == 1 and text == "附件":
                doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
            if level == 2 and text.startswith("附件"):
                doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
            h = doc.add_heading(level=min(level, 4))
            para_format(h, first_indent=level >= 2, line=Pt(28), before=6 if level == 1 else 0)
            font = {1: FONT_HEI, 2: FONT_KAI}.get(level, FONT_FANG)
            add_inline(h, text, font, SIZE_BODY, bold=level in (1, 3))
            i += 1
            continue
        if re.match(r"^[-*]\s+", s):
            p = doc.add_paragraph()
            para_format(p)
            add_inline(p, "· " + re.sub(r"^[-*]\s+", "", s), FONT_FANG, SIZE_BODY)
            i += 1
            continue
        if s.startswith(">"):
            p = doc.add_paragraph()
            para_format(p)
            add_inline(p, s.lstrip("> "), FONT_KAI, SIZE_BODY)
            i += 1
            continue
        p = doc.add_paragraph()
        para_format(p)
        add_inline(p, s, FONT_FANG, SIZE_BODY)
        i += 1

    doc.save(dst)


if __name__ == "__main__":
    convert(sys.argv[1], sys.argv[2])
